"""
Migra "Análisis Mercari.xlsx" → ../datos-privados/02_datos_excel.sql  +  ../datos-privados/fotos_migracion/<id>.<ext>

Uso:  python3 herramientas/migrar_excel.py "Análisis Mercari.xlsx"
Requiere: pip install openpyxl
"""
import sys, os, datetime, re
import openpyxl

ORIGEN = sys.argv[1] if len(sys.argv) > 1 else "Análisis Mercari.xlsx"
# Los datos van FUERA del repo (carpeta hermana "datos-privados") para no publicarlos en GitHub
RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PRIV = os.path.join(RAIZ, "datos-privados")
SALIDA_SQL = os.path.join(PRIV, "02_datos_excel.sql")
DIR_FOTOS = os.path.join(PRIV, "fotos_migracion")

ESTATUS = {"SOLD": "VENDIDO", "CHECK": "RECIBIDO", "TRANSIT": "TRANSITO"}
MARCAS_ALIAS = {"DIOR": "CHRISTIAN DIOR", "ZIMMERMAN": "ZIMMERMANN", "AQUAZZURA": "AQUAZZURA"}


def q(v):
    """Literal SQL."""
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return "null"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return repr(round(float(v), 4)).rstrip("0").rstrip(".") if isinstance(v, float) else str(v)
    if isinstance(v, (datetime.datetime, datetime.date)):
        return "'" + v.strftime("%Y-%m-%d") + "'"
    return "'" + str(v).strip().replace("'", "''") + "'"


def num(v):
    if v in (None, "", "TBD"):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def fecha(v):
    return v if isinstance(v, (datetime.datetime, datetime.date)) else None


def marca_norm(m):
    if not m:
        return None
    k = re.sub(r"\s+", " ", str(m)).strip().upper()
    return MARCAS_ALIAS.get(k, k)


wb = openpyxl.load_workbook(ORIGEN, data_only=True)
out = ["-- Generado automáticamente desde el Excel. Ejecutar DESPUÉS de 01_esquema.sql.",
       "-- Se puede correr más de una vez: no duplica registros.", "begin;", ""]
avisos = []

# ---------------- Main ----------------
ws = wb["Main"]
articulos = []
for idx, r in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
    if not r[1] or not r[2]:
        continue
    articulos.append((idx, r))
ids = {str(r[1]).strip() for _, r in articulos}

# ---------------- Shipping → paquetes ----------------
paquetes = {}
for r in wb["Shipping"].iter_rows(min_row=4, values_only=True):
    n = (r[1] or "").strip() if isinstance(r[1], str) else r[1]
    if not n:
        continue
    if n in paquetes:
        avisos.append(f"Paquete {n} venía duplicado en la hoja Shipping (se dejó uno).")
        continue
    peso = num(r[13])
    if peso is not None and (r[14] or "g").lower().startswith("k"):
        peso *= 1000
    paquetes[n] = dict(numero=n, total_jpy=num(r[2]), envio_intl_jpy=num(r[3]), photo_service_jpy=num(r[4]),
                       otros_jpy=num(r[5]), total_usd=num(r[6]), total_mxn=num(r[7]), tipo_envio=r[8],
                       seguro=bool(r[9]), inspeccion=bool(r[10]), empaque_protector=bool(r[11]),
                       dimensiones=r[12], peso_g=peso)
for _, r in articulos:
    n = r[9]
    if n and n != "TBD" and n not in paquetes:
        paquetes[n] = dict(numero=n, notas="Sin detalle en hoja Shipping")

out.append("-- Paquetes")
for p in paquetes.values():
    cols = list(p.keys())
    out.append(f"insert into paquetes ({', '.join(cols)}) values ({', '.join(q(p[c]) for c in cols)}) on conflict (numero) do nothing;")

# ---------------- Marcas ----------------
marcas = sorted({marca_norm(r[3]) for _, r in articulos if r[3]})
out.append("\n-- Marcas")
for m in marcas:
    out.append(f"insert into marcas (nombre) values ({q(m)}) on conflict (nombre) do nothing;")

# ---------------- Ventas (hoja) ----------------
ventas_hoja = {}
sin_articulo = []
for idx, r in enumerate(wb["Ventas"].iter_rows(min_row=3, values_only=True), start=3):
    vid = (r[1] or "").strip() if isinstance(r[1], str) else r[1]
    if not any(r[1:16]):
        continue
    if not vid:
        sin_articulo.append(idx)
        continue
    ventas_hoja[vid] = r
envios_hoja = {r[1]: num(r[2]) for r in wb["Envíos"].iter_rows(min_row=3, values_only=True) if r[1]}

# ---------------- Artículos + ventas ----------------
out.append("\n-- Artículos")
ventas_sql = []
sin_estatus = 0
for idx, r in articulos:
    aid = str(r[1]).strip()
    st = ESTATUS.get((r[16] or "").strip().upper()) if r[16] else None
    if not st:
        st = "COMPRADO"
        sin_estatus += 1
    w, x = num(r[22]), num(r[23])
    vendido = st == "VENDIDO"
    paq = r[9] if r[9] and r[9] != "TBD" else None
    marca = marca_norm(r[3])
    valores = {
        "id": aid, "nombre": str(r[2]).strip(), "color": r[4], "vendedor": r[5], "pagado_por": r[6],
        "orden_compra": r[7], "rastreo_correo": r[8], "paquete_numero": paq,
        "precio_jpy": num(r[11]), "precio_usd": num(r[12]), "pago_mxn": num(r[13]), "fecha_compra": fecha(r[14]),
        "costo_envio_mxn": num(r[18]), "impuestos_mxn": num(r[20]) or 0,
        "precio_objetivo": w, "estatus": st,
        "canal_objetivo": "GoTrendier" if (not vendido and w and x and x > 19) else None,
        "notas": f"Migrado de Excel (fila {idx})",
    }
    cols = list(valores.keys())
    vals = [q(valores[c]) for c in cols]
    cols.append("marca_id")
    vals.append(f"(select id from marcas where nombre = {q(marca)})" if marca else "null")
    out.append(f"insert into articulos ({', '.join(cols)}) values ({', '.join(vals)}) on conflict (id) do nothing;")

    if num(r[12]) and num(r[11]) and num(r[12]) > num(r[11]) / 20:
        avisos.append(f"{aid}: precio USD ({r[12]}) parece incorrecto para {r[11]} JPY — revísalo.")

    if vendido:
        vh = ventas_hoja.get(aid)
        canal = None
        if (x or 0) == 0:
            canal = "Directo"
        if vh and vh[9]:
            canal = str(vh[9]).strip()
        venta = {
            "articulo_id": aid, "precio_venta": w or 0, "comision_mxn": x or 0, "canal": canal,
            "envio_mxn": envios_hoja.get(aid) or 0, "fecha_venta": None,
            "notas": "Migrado de Excel",
        }
        if vh:
            notas_extra = [str(v) for v in (vh[13], vh[14]) if v]
            venta.update(fecha_venta=fecha(vh[2]), fecha_envio=fecha(vh[3]), destino=vh[4], cliente=vh[5],
                         telefono=str(int(vh[6])) if isinstance(vh[6], float) else vh[6], correo=vh[7],
                         paqueteria=vh[8], monto_depositado=num(vh[10]), fecha_deposito=fecha(vh[11]),
                         cuenta_deposito=vh[12], folio=vh[15],
                         notas="Migrado de Excel" + (" · " + " / ".join(notas_extra) if notas_extra else ""))
        if not w:
            avisos.append(f"{aid}: vendido sin precio de venta en el Excel — completa la venta.")
        c = list(venta.keys())
        ventas_sql.append(f"insert into ventas ({', '.join(c)}) values ({', '.join(q(venta[k]) for k in c)}) on conflict (articulo_id) do nothing;")

for vid in ventas_hoja:
    if vid not in ids:
        avisos.append(f"Hoja Ventas: el artículo {vid} no existe en Main.")
if sin_articulo:
    avisos.append(f"Hoja Ventas: {len(sin_articulo)} depósitos sin ID de artículo (filas {', '.join(map(str, sin_articulo))}) no se migraron; regístralos en la venta que corresponda.")

out.append("\n-- Ventas")
out += ventas_sql

# ---------------- Cargos TC ----------------
out.append("\n-- Cargos a tarjeta")
ws = wb["Registros Cargo TC"]
n_cargos = 0
for idx, r in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
    for lado, (o, art, conc, monto, f_cargo, tarj, banco, pago, edo, fpago) in (
        ("izq", (r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8], r[9], r[10])),
        ("der", (r[16], None, r[17], r[18], r[19], r[20], r[21], r[22], r[23], r[24])),
    ):
        if monto in (None, ""):
            continue
        personal = lado == "der"
        art = str(art).strip() if art else None
        nota = f"Excel: Registros Cargo TC fila {idx} ({lado})"
        if art and art not in ids:
            nota += f" · artículo {art} no encontrado"
            art = None
        concepto = conc if conc in ("Compra", "Envío", "Impuestos") else "Otro"
        valores = dict(fecha=fecha(f_cargo), articulo_id=art, concepto=concepto,
                       descripcion=None if concepto != "Otro" else conc, monto_mxn=num(monto),
                       tarjeta=tarj, banco=banco, estado="PAGADO" if (edo or "").upper() == "PAGADO" else "PENDIENTE",
                       fecha_pago=fecha(fpago), monto_pagado=num(pago), es_personal=personal, notas=nota)
        c = list(valores.keys())
        out.append(f"insert into cargos_tc ({', '.join(c)}) select {', '.join(q(valores[k]) for k in c)} "
                   f"where not exists (select 1 from cargos_tc where notas = {q(nota)});")
        n_cargos += 1

out += ["", "commit;"]
os.makedirs(os.path.dirname(SALIDA_SQL), exist_ok=True)
with open(SALIDA_SQL, "w", encoding="utf-8") as f:
    f.write("\n".join(out) + "\n")

# ---------------- Fotos ----------------
os.makedirs(DIR_FOTOS, exist_ok=True)
ws = wb["Main"]
fila_a_id = {idx: str(r[1]).strip() for idx, r in articulos}
n_fotos = 0
for im in openpyxl.load_workbook(ORIGEN)["Main"]._images:
    fila = im.anchor._from.row + 1
    aid = fila_a_id.get(fila)
    if not aid:
        continue
    data = im._data()
    ext = "png" if data[:4] == b"\x89PNG" else "jpg"
    with open(os.path.join(DIR_FOTOS, f"{aid}.{ext}"), "wb") as f:
        f.write(data)
    n_fotos += 1

print(f"Artículos: {len(articulos)} ({sin_estatus} sin estatus → COMPRADO)")
print(f"Paquetes: {len(paquetes)} · Marcas: {len(marcas)} · Ventas: {len(ventas_sql)} · Cargos: {n_cargos} · Fotos: {n_fotos}")
for a in avisos:
    print(" ⚠", a)
